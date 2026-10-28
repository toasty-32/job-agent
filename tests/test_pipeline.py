import shutil
from pathlib import Path
from types import SimpleNamespace

import docx

from job_agent.config import Preferences
from job_agent.documents import markdown_to_docx
from job_agent.knowledge_base import load_knowledge_base, render_corpus
from job_agent.linkedin import Job
from job_agent.llm import Application, CareerModel, FitScore
from job_agent.pipeline import JobAgent
from job_agent.store import filter_reason

ROOT = Path(__file__).resolve().parent.parent


def make_prefs(tmp_path) -> Path:
    path = tmp_path / "preferences.yaml"
    shutil.copy(ROOT / "preferences.example.yaml", path)
    return path


def test_example_preferences_load(tmp_path):
    prefs = Preferences.load(make_prefs(tmp_path))
    assert prefs.search.keywords
    assert prefs.claude.model == "claude-opus-5-5"
    assert "priorities" in prefs.raw_yaml


def test_knowledge_base_loads_text_and_docx(tmp_path):
    (tmp_path / "README.md").write_text("ignored")
    (tmp_path / "_draft.md").write_text("ignored")
    (tmp_path / "profile.md").write_text("# Me\nPython")
    (tmp_path / "skills.json").write_text('{"skills": ["R"]}')
    d = docx.Document()
    d.add_paragraph("Old CV line")
    (tmp_path / "sub").mkdir()
    d.save(str(tmp_path / "sub" / "cv.docx"))
    files = load_knowledge_base(tmp_path)
    assert [f.path for f in files] == ["profile.md", "skills.json", "sub/cv.docx"]
    corpus = render_corpus(files)
    assert '<document path="sub/cv.docx">\nOld CV line' in corpus


def test_filters():
    prefs = Preferences.load(ROOT / "preferences.example.yaml")
    assert filter_reason(Job("1", "Bioinformatics Intern", "X", "", ""), prefs.filters)
    assert filter_reason(Job("2", "Scientist", "X", "", ""), prefs.filters) is None


def test_markdown_to_docx(tmp_path):
    path = tmp_path / "cv.docx"
    markdown_to_docx("# Alex\n\nalex@x.com\n\n## Experience\n- Built **fast** pipelines\n- Led team", path)
    paragraphs = docx.Document(str(path)).paragraphs
    assert [p.style.name for p in paragraphs] == ["Title", "Normal", "Heading 1", "List Bullet", "List Bullet"]
    assert paragraphs[3].runs[1].bold and paragraphs[3].text == "Built fast pipelines"


class FakeStream:
    def __init__(self, parsed):
        self.message = SimpleNamespace(stop_reason="end_turn", parsed_output=parsed)

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def get_final_message(self):
        return self.message


class FakeClient:
    def __init__(self):
        self.calls = []
        self.beta = SimpleNamespace(messages=SimpleNamespace(stream=self.stream))

    def stream(self, **kwargs):
        self.calls.append(kwargs)
        if kwargs["output_format"] is FitScore:
            return FakeStream(FitScore(score=120, verdict="strong", summary="Great", matching_strengths=["a"], gaps=["b"], dealbreakers=[]))
        return FakeStream(Application(
            cv_markdown="# Alex Example\n\n## Skills\n- Python",
            cover_letter_markdown="Dear hiring team,\n\nHello.",
            tailoring_notes=["Led with pipelines"],
            unsupported_requirements=["Kubernetes"],
        ))


def test_score_and_write_end_to_end(tmp_path):
    prefs = Preferences.load(make_prefs(tmp_path))
    agent = JobAgent(prefs, ROOT / "knowledge_base", tmp_path / "out", log=lambda _: None)
    client = FakeClient()
    agent._model = CareerModel(prefs, load_knowledge_base(ROOT / "knowledge_base"), client=client)

    job = agent.add_job("https://acme.example/jobs/9", "We need Python and Kubernetes.", "Data Scientist", "Acme")
    assert job.job_id.startswith("manual-")
    assert agent.add_job("https://acme.example/jobs/9", "We need Python and Kubernetes.", "Data Scientist", "Acme").job_id == job.job_id

    assert agent.score(job.job_id).score == 100  # clamped
    _, paths = agent.write(job.job_id)

    folder = Path(paths["folder"])
    assert {p.name for p in folder.iterdir()} == {"cv.md", "cv.docx", "cover_letter.md", "cover_letter.docx", "notes.md"}
    assert "Kubernetes" in (folder / "notes.md").read_text()

    # Both calls share the same cached system prompt carrying the knowledge base and preferences.
    assert client.calls[0]["system"] == client.calls[1]["system"]
    system_text = client.calls[0]["system"][0]["text"]
    assert "Alex Example" in system_text and "<preferences>" in system_text
    assert client.calls[0]["output_config"] == {"effort": "low"}
    assert client.calls[1]["output_config"] == {"effort": "high"}
    assert client.calls[0]["fallbacks"] == "default"

    reloaded = JobAgent(prefs, ROOT / "knowledge_base", tmp_path / "out", log=lambda _: None)
    assert reloaded.store.ranked()[0][1]["documents"]["folder"] == paths["folder"]
