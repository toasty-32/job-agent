"""Writes generated CVs and cover letters to disk as Markdown and Word (.docx)."""
from __future__ import annotations

import re
from pathlib import Path

import docx
from docx.shared import Pt

from .linkedin import Job
from .llm import Application, FitScore

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\))")


def slugify(text: str, limit: int = 40) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return slug[:limit].strip("-") or "untitled"


def _add_inline(paragraph, text: str) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        elif part.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            paragraph.add_run(label if label == url else f"{label} ({url})")
        else:
            paragraph.add_run(part)


def markdown_to_docx(markdown: str, path: Path) -> None:
    """Renders the Markdown subset the writer is asked to use: headings, bullets, paragraphs."""
    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for section in document.sections:
        section.left_margin = section.right_margin = Pt(54)
        section.top_margin = section.bottom_margin = Pt(48)

    paragraph_lines: list[str] = []

    def flush() -> None:
        if paragraph_lines:
            _add_inline(document.add_paragraph(), " ".join(paragraph_lines))
            paragraph_lines.clear()

    for raw in markdown.splitlines():
        line = raw.rstrip()
        heading = re.match(r"^(#{1,4})\s+(.*)", line)
        bullet = re.match(r"^\s*[-*+]\s+(.*)", line)
        if not line.strip():
            flush()
        elif heading:
            flush()
            document.add_heading(heading.group(2).strip(), level=len(heading.group(1)) - 1 or 0)
        elif bullet:
            flush()
            _add_inline(document.add_paragraph(style="List Bullet"), bullet.group(1))
        elif re.match(r"^\s*(-{3,}|\*{3,})\s*$", line):
            flush()
        else:
            paragraph_lines.append(line.strip())
    flush()
    document.save(str(path))


def save_application(output_dir: Path, job: Job, application: Application, score: FitScore | None) -> dict:
    folder = output_dir / "applications" / f"{slugify(job.company, 25)}--{slugify(job.title)}--{job.job_id}"
    folder.mkdir(parents=True, exist_ok=True)

    (folder / "cv.md").write_text(application.cv_markdown, encoding="utf-8")
    (folder / "cover_letter.md").write_text(application.cover_letter_markdown, encoding="utf-8")
    markdown_to_docx(application.cv_markdown, folder / "cv.docx")
    markdown_to_docx(application.cover_letter_markdown, folder / "cover_letter.docx")

    notes = [f"# {job.title} — {job.company}", "", f"- Link: {job.url}", f"- Location: {job.location}"]
    if score:
        notes += [f"- Fit score: {score.score}/100 ({score.verdict})", "", score.summary]
        if score.gaps:
            notes += ["", "## Gaps", *[f"- {g}" for g in score.gaps]]
    notes += ["", "## How the documents were tailored", *[f"- {n}" for n in application.tailoring_notes]]
    if application.unsupported_requirements:
        notes += [
            "", "## Requirements your knowledge base has no evidence for",
            "Add evidence to the knowledge base if you do have these, then regenerate.",
            *[f"- {r}" for r in application.unsupported_requirements],
        ]
    notes += ["", "## Job description", "", job.description]
    (folder / "notes.md").write_text("\n".join(notes) + "\n", encoding="utf-8")
    return {"folder": str(folder)}
