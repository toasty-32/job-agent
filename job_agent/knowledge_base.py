"""Reads every supported file in the knowledge base folder into one text corpus."""
from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path

TEXT_SUFFIXES = {".md", ".markdown", ".txt", ".yaml", ".yml", ".json"}
DOC_SUFFIXES = {".docx", ".pdf"}


@dataclass
class KnowledgeFile:
    path: str  # relative to the knowledge base root
    text: str


def _read_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(line for line in lines if line.strip())


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)


def _read(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".json":
        return json.dumps(json.loads(path.read_text(encoding="utf-8")), indent=2, ensure_ascii=False)
    return path.read_text(encoding="utf-8")


def load_knowledge_base(kb_dir: Path) -> list[KnowledgeFile]:
    if not kb_dir.is_dir():
        raise FileNotFoundError(f"Knowledge base folder {kb_dir} does not exist")
    files = []
    # Sorted so the corpus is byte-identical between runs, which keeps the prompt cache warm.
    for path in sorted(kb_dir.rglob("*")):
        rel = path.relative_to(kb_dir)
        if not path.is_file() or any(part.startswith((".", "_")) for part in rel.parts):
            continue
        if rel.as_posix() == "README.md" or path.suffix.lower() not in TEXT_SUFFIXES | DOC_SUFFIXES:
            continue
        text = _read(path).strip()
        if text:
            files.append(KnowledgeFile(rel.as_posix(), text))
    if not files:
        raise ValueError(f"Knowledge base {kb_dir} has no readable files — add your profile first")
    return files


def render_corpus(files: list[KnowledgeFile]) -> str:
    return "\n\n".join(
        f'<document path="{f.path}">\n{f.text}\n</document>' for f in files
    )
